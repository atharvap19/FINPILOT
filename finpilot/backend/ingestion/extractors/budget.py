"""Budget sheets: what someone *planned* to spend, by category.

A budget is not a ledger. Its rows are categories rather than payments, it has
no dates, and its "actual" column is a total of money that is already itemised
in the statements. Importing those rows as transactions would invent dates and
count the same spending twice, so budgets live in their own table -- and that
is exactly what makes them useful later, because budget-versus-actual is a
comparison the insights stage can only make if both sides exist separately.

Recognising one is a matter of shape: a category column, at least one amount
column, and *no* date column anywhere. That last part is what separates a
budget from an expense sheet, which :mod:`ingestion.extractors.tabular` handles
as ordinary transactions.

The variance column is read but never trusted -- it is recomputed here, in
code, from the two amounts (principle 2). A sheet whose own arithmetic is wrong
is reported rather than silently corrected.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from pydantic import BaseModel, ConfigDict, field_validator

from ingestion.mapping import normalize_header
from ingestion.normalize import clean_text, is_blank, parse_amount, parse_date, quantize_money

# --------------------------------------------------------------------------
# recognising the columns
# --------------------------------------------------------------------------

_CATEGORY = ("category", "categories", "head", "heads", "item", "items", "particulars",
             "particular", "expense", "expenses", "expense head", "description", "type",
             "bucket", "segment")
_BUDGETED = ("budget", "budgeted", "monthly budget", "planned", "plan", "target",
             "allocated", "allocation", "limit", "estimate", "estimated", "budget inr",
             "monthly budget inr")
_ACTUAL = ("actual", "actuals", "actual spend", "spent", "spend", "used", "real",
           "incurred", "actual inr", "actual spend inr")
_VARIANCE = ("variance", "difference", "diff", "delta", "remaining", "left", "balance",
             "over under", "saving", "savings")
_NOTES = ("notes", "note", "remark", "remarks", "comment", "comments", "memo")
_PERIOD = ("period", "month", "billing period", "for the month", "cycle")

#: Rows that summarise the sheet rather than describe a category.
_TOTAL_ROW = re.compile(r"^(total|grand total|sum|overall|net)\b", re.IGNORECASE)


def _match(header: str, names: tuple[str, ...]) -> bool:
    norm = normalize_header(header)
    return any(norm == n or norm.startswith(f"{n} ") or n in norm for n in names)


@dataclass
class BudgetColumns:
    category: int | None = None
    budgeted: int | None = None
    actual: int | None = None
    variance: int | None = None
    notes: int | None = None
    period: int | None = None

    @property
    def usable(self) -> bool:
        """A budget needs a category and at least one amount to mean anything."""
        return self.category is not None and (
            self.budgeted is not None or self.actual is not None
        )


def find_budget_columns(headers: list[str]) -> BudgetColumns:
    columns = BudgetColumns()
    for index, header in enumerate(headers):
        if not clean_text(header):
            continue
        # most specific first: "actual spend" must not be claimed by "spend"
        # under a looser rule, and "monthly budget" must not become a period
        if columns.actual is None and _match(header, _ACTUAL):
            columns.actual = index
        elif columns.budgeted is None and _match(header, _BUDGETED):
            columns.budgeted = index
        elif columns.variance is None and _match(header, _VARIANCE):
            columns.variance = index
        elif columns.category is None and _match(header, _CATEGORY):
            columns.category = index
        elif columns.notes is None and _match(header, _NOTES):
            columns.notes = index
        elif columns.period is None and _match(header, _PERIOD):
            columns.period = index
    return columns


# --------------------------------------------------------------------------
# the extracted shape
# --------------------------------------------------------------------------


class BudgetLine(BaseModel):
    """One category's plan, and what was actually spent against it."""

    model_config = ConfigDict(extra="forbid")

    category: str
    budgeted: Decimal | None = None
    actual: Decimal | None = None
    notes: str | None = None
    source_ref: str

    @field_validator("budgeted", "actual", mode="before")
    @classmethod
    def _money(cls, v):
        if v in (None, ""):
            return None
        if isinstance(v, (int, float, Decimal)):
            return quantize_money(v)
        parsed = parse_amount(v)
        return None if parsed is None else parsed

    @property
    def variance(self) -> Decimal | None:
        """Budgeted minus actual. Always recomputed, never read off the sheet.

        Positive means under budget. The sheet's own variance column is only
        used to check this, in :func:`check_variance`.
        """
        if self.budgeted is None or self.actual is None:
            return None
        return quantize_money(self.budgeted - self.actual)

    @property
    def over_budget(self) -> bool:
        variance = self.variance
        return variance is not None and variance < 0


@dataclass
class BudgetRecord:
    """A budget sheet, parsed."""

    period: str | None
    lines: list[BudgetLine]
    doc_id: str
    source_ref: str
    issues: list[str] = field(default_factory=list)

    @property
    def total_budgeted(self) -> Decimal:
        return quantize_money(
            sum((line.budgeted for line in self.lines if line.budgeted), Decimal("0"))
        )

    @property
    def total_actual(self) -> Decimal:
        return quantize_money(
            sum((line.actual for line in self.lines if line.actual), Decimal("0"))
        )


# --------------------------------------------------------------------------
# parsing
# --------------------------------------------------------------------------

#: Tolerance when checking the sheet's own variance column against ours.
VARIANCE_TOLERANCE = Decimal("0.05")


def check_variance(line: BudgetLine, stated: Decimal | None) -> str | None:
    """Does the sheet's own variance agree with the arithmetic?"""
    if stated is None:
        return None
    ours = line.variance
    if ours is None:
        return None
    if abs(ours - stated) <= VARIANCE_TOLERANCE:
        return None
    return (
        f"{line.category}: the sheet says the variance is {stated:.2f}, but "
        f"{(line.budgeted or 0):.2f} budgeted less {(line.actual or 0):.2f} spent "
        f"is {ours:.2f}"
    )


def parse_budget_rows(
    headers: list[str],
    rows: list[list[str]],
    *,
    filename: str,
    sheet_name: str | None = None,
    doc_id: str = "",
) -> BudgetRecord:
    """Turn a recognised budget table into a :class:`BudgetRecord`."""
    columns = find_budget_columns(headers)
    lines: list[BudgetLine] = []
    issues: list[str] = []
    period_values: list[str] = []

    for position, row in enumerate(rows, start=1):
        def cell(index: int | None) -> str:
            return clean_text(row[index]) if index is not None and index < len(row) else ""

        category = cell(columns.category)
        if not category or _TOTAL_ROW.match(category):
            continue  # a totals row is the sheet's arithmetic, not a category
        budgeted, actual = cell(columns.budgeted), cell(columns.actual)
        if is_blank(budgeted) and is_blank(actual):
            continue

        line = BudgetLine(
            category=category,
            budgeted=budgeted or None,
            actual=actual or None,
            notes=cell(columns.notes) or None,
            source_ref=f"{filename}#{position}",
        )
        lines.append(line)

        problem = check_variance(line, parse_amount(cell(columns.variance)))
        if problem:
            issues.append(problem)
        if (value := cell(columns.period)):
            period_values.append(value)

    period = period_values[0] if period_values else _period_from(filename, sheet_name)
    return BudgetRecord(
        period=period,
        lines=lines,
        doc_id=doc_id,
        source_ref=f"{filename}#1",
        issues=issues,
    )


_MONTHS = ("january", "february", "march", "april", "may", "june", "july", "august",
           "september", "october", "november", "december")
_MONTH_IN_NAME = re.compile(
    r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\b[_\-\s]*(\d{4})?",
    re.IGNORECASE,
)


def _period_from(filename: str, sheet_name: str | None) -> str | None:
    """``monthly_budget_aug_2026.xlsx`` -> ``Aug 2026``.

    The sheet tab is tried first: someone who named it "August Budget" was
    being more deliberate than whoever named the file.
    """
    for candidate in (sheet_name or "", Path(filename).stem.replace("_", " ")):
        if not candidate:
            continue
        parsed = parse_date(candidate)
        if parsed is not None:
            return parsed.strftime("%b %Y")
        match = _MONTH_IN_NAME.search(candidate)
        if match:
            month = match.group(1).lower()[:3].title()
            year = match.group(2)
            return f"{month} {year}" if year else month
        for month in _MONTHS:
            if month in candidate.lower():
                return month.title()
    return None


def looks_like_a_budget(headers: list[str], rows: list[list[str]]) -> bool:
    """A category column, an amount column, and no dates anywhere.

    The absence of dates is the decisive part. An expense sheet with a date
    column is a list of transactions and belongs in the ledger; a sheet of
    category totals with no dates is a plan.
    """
    if not find_budget_columns(headers).usable:
        return False
    for column in range(len(headers)):
        values = [r[column] for r in rows if column < len(r) and not is_blank(r[column])]
        if not values:
            continue
        dated = sum(1 for v in values if parse_date(v) is not None)
        if dated / len(values) >= 0.6:
            return False
    return True


def extract_budget(
    path: Path, file_type: str, *, filename: str | None = None, doc_id: str = ""
) -> BudgetRecord | None:
    """Read a budget sheet. ``None`` when the file is not one."""
    from ingestion.extractors.tabular import _best_shaped_row, read_grid

    name = filename or path.name
    grid = read_grid(path, file_type)
    if not grid:
        return None
    header_index = _best_shaped_row(grid)
    if header_index is None:
        return None

    headers = [clean_text(c) for c in grid[header_index]]
    rows = [
        [clean_text(c) for c in r]
        for r in grid[header_index + 1 :]
        if any(not is_blank(c) for c in r)
    ]
    if not looks_like_a_budget(headers, rows):
        return None

    record = parse_budget_rows(
        headers, rows, filename=name, sheet_name=_sheet_name(path, file_type), doc_id=doc_id
    )
    return record if record.lines else None


def _sheet_name(path: Path, file_type: str) -> str | None:
    from ingestion.detect import XLS, XLSX

    if file_type not in (XLSX, XLS):
        return None
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        try:
            return wb.sheetnames[0] if wb.sheetnames else None
        finally:
            wb.close()
    except Exception:  # noqa: BLE001
        return None


def budget_to_rows(record: BudgetRecord, *, as_of: dt.date | None = None) -> list:
    """Flatten a record into :class:`~ingestion.schema.Budget` rows."""
    from ingestion.schema import Budget, make_budget_id

    out = []
    for line in record.lines:
        out.append(
            Budget(
                budget_id=make_budget_id(period=record.period, category=line.category),
                period=record.period,
                category=line.category,
                budgeted_amount=line.budgeted,
                actual_amount=line.actual,
                variance=line.variance,
                notes=line.notes,
                source_ref=line.source_ref,
            )
        )
    return out
